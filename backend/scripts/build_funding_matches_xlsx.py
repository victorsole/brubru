#!/usr/bin/env python3.12
"""One Excel workbook per client from build_funding_matches.py's --json output.

The 15 September 2026 workbooks were made by a scratch script that was never
saved, so every three-day run would have had to rebuild the layout from memory.
This is that layout, kept: a "Matches" sheet (title, subtitle, frozen header,
filters, hyperlinked titles, High/Medium badges) and a "How to read this" sheet.

Rows: High and Medium matches from the main run and the portal-verified gap run
(`<slug>` and `<slug>_gap` in the JSON), de-duplicated by URL, deadline at least
two days after the report date (nothing that closes before a client can act),
sorted High first, then by deadline.

Usage (from the repo root):
    python3.12 backend/scripts/build_funding_matches.py --date 2026-09-23 --json /tmp/fm.json
    python3.12 backend/scripts/build_funding_matches_xlsx.py --json /tmp/fm.json --date 2026-09-23
"""
from __future__ import annotations

import argparse
import datetime as dt
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

REPO = Path(__file__).resolve().parents[2]

ORGS = {
    "terraqui": ("Terraqui", "Estudi Jurídic Ambiental, S.L.P. (Terraqui), Barcelona"),
    "gbsb": ("GBSB", "GBSB Global Business School"),
    "cadence": ("Cadence", "Cadence Design Systems"),
}
BLUE, GREEN, AMBER, GREY, LINK = "1F3A8A", "15803D", "B45309", "555555", "1D4ED8"
HEADERS = ["#", "Match", "Opportunity (click to open)", "Programme or buyer", "Source", "Deadline",
           "Days left", "Why it fits", "Realistic role", "Check before applying", "Link"]
WIDTHS = [5, 10, 55, 32, 24, 12, 9, 70, 38, 50, 40]


def _long(d: dt.date) -> str:
    return f"{d.day} {d.strftime('%B %Y')}"


def _source(raw: str | None) -> str:
    raw = raw or ""
    if raw.startswith("EU Funding & Tenders"):
        return "EU Funding & Tenders Portal"
    if raw.startswith("TED"):
        return "TED (EU public procurement)"
    return raw or "EU source"


def rows_for(data: dict, slug: str, earliest: dt.date) -> list[dict]:
    seen, out = set(), []
    for m in (data.get(slug) or []) + (data.get(f"{slug}_gap") or []):
        if m.get("strength") not in ("High", "Medium") or m.get("url") in seen:
            continue
        deadline = dt.date.fromisoformat(m["deadline"][:10])
        if deadline < earliest:
            continue
        seen.add(m["url"])
        why, _, role = (m.get("why") or "").partition(" Realistic role: ")
        out.append({**m, "deadline_d": deadline, "why_text": why.strip(),
                    "role": role.strip().rstrip("."), "check": "\n".join(f"- {c}" for c in (m.get("caveats") or []))})
    out.sort(key=lambda r: (r["strength"] != "High", r["deadline_d"], r["title"]))
    return out


def write(slug: str, rows: list[dict], report: dt.date, earliest: dt.date, out_dir: Path) -> Path:
    short, full = ORGS[slug]
    wb = Workbook()
    ws = wb.active
    ws.title = "Matches"
    ws["A1"] = f"EU funding and procurement opportunities for {full}"
    ws["A1"].font = Font(bold=True, size=14, color=BLUE)
    ws["A2"] = (f"Prepared by Brubru (brubru.beresol.eu) on {_long(report)}. High and medium matches only; "
                f"open or forthcoming, deadline on or after {_long(earliest)}. "
                "Always read the official call or notice before applying.")
    ws["A2"].font = Font(italic=True, size=10, color=GREY)
    for col, (h, w) in enumerate(zip(HEADERS, WIDTHS), start=1):
        c = ws.cell(row=4, column=col, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=BLUE)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[c.column_letter].width = w
    top = Alignment(vertical="top", wrap_text=True)
    for i, r in enumerate(rows, start=1):
        row = 4 + i
        vals = [i, r["strength"], r["title"], r.get("programme") or "", _source(r.get("source")),
                r["deadline_d"], (r["deadline_d"] - report).days, r["why_text"], r["role"], r["check"], r["url"]]
        for col, v in enumerate(vals, start=1):
            c = ws.cell(row=row, column=col, value=v)
            c.alignment = top
        badge = ws.cell(row=row, column=2)
        badge.font = Font(bold=True, color="FFFFFF")
        badge.fill = PatternFill("solid", fgColor=GREEN if r["strength"] == "High" else AMBER)
        badge.alignment = Alignment(horizontal="center", vertical="top")
        ws.cell(row=row, column=6).number_format = "dd/mm/yyyy"
        for col in (3, 11):
            c = ws.cell(row=row, column=col)
            c.hyperlink = r["url"]
            c.font = Font(color=LINK, underline="single", size=11 if col == 3 else 9)
    ws.freeze_panes = "D5"
    ws.auto_filter.ref = f"A4:K{4 + max(len(rows), 1)}"

    how = wb.create_sheet("How to read this")
    lines = [
        "How to read this file", None,
        "Match: High means the opportunity fits the organisation's core activity and it could realistically apply "
        "or bid; Medium means a good fit that needs a partner, a consortium or a closer eligibility check.",
        "Source: EU Funding & Tenders Portal calls (Horizon Europe, LIFE, Digital Europe, Erasmus+, and others), EU "
        "public procurement notices published on TED, and procurement pages of EU institutions and agencies.",
        f"Deadline: the submission deadline shown on the official page on {_long(report)}. Multi-stage calls show the "
        "next deadline. Dates can change: the official page prevails.",
        "Why it fits / Realistic role / Check before applying: Brubru's assessment against the organisation's "
        "published activities. It is a starting point for a decision, not legal or grant advice.",
        "TED notices: many public tenders are published in the buyer's language; the title is kept as published.",
        None,
        "Brubru, All the EU, with AI. https://brubru.beresol.eu/  ·  hello@beresol.eu",
    ]
    for n, text in enumerate(lines, start=1):
        c = how.cell(row=n, column=1, value=text)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if n == 1:
            c.font = Font(bold=True, size=14, color=BLUE)
    how.column_dimensions["A"].width = 120

    path = out_dir / f"Brubru_EU_funding_matches_{short}_{report.isoformat()}.xlsx"
    wb.save(path)
    return path


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--json", required=True, help="the --json file written by build_funding_matches.py")
    ap.add_argument("--date", default=dt.date.today().isoformat())
    ap.add_argument("--out-dir", default=str(REPO / "docs" / "funding_matches"))
    ap.add_argument("--only", nargs="*", default=list(ORGS), help="slugs to build")
    a = ap.parse_args()
    report = dt.date.fromisoformat(a.date)
    earliest = report + dt.timedelta(days=2)
    data = json.loads(Path(a.json).read_text())
    out_dir = Path(a.out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    for slug in a.only:
        rows = rows_for(data, slug, earliest)
        path = write(slug, rows, report, earliest, out_dir)
        hi = sum(r["strength"] == "High" for r in rows)
        print(f"[OK] {path.name}: {len(rows)} rows ({hi} High, {len(rows) - hi} Medium)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
